import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# ── Task 1: Confidence Interval ───────────────────────────────────────────────
p = 0.45
n = 8000
z = 1.96

se_1 = math.sqrt(p * (1 - p) / n)
me_1 = z * se_1
ci_lower = p - me_1
ci_upper = p + me_1

# ── Task 2: A/B Test ──────────────────────────────────────────────────────────
p_a, n_a = 0.55, 12000
p_b, n_b = 0.58, 12000

p_pool = (n_a * p_a + n_b * p_b) / (n_a + n_b)
se_2 = math.sqrt(p_pool * (1 - p_pool) * (1 / n_a + 1 / n_b))
z_stat = (p_b - p_a) / se_2
z_crit = 1.96
is_significant = abs(z_stat) > z_crit

# ── Style helpers ─────────────────────────────────────────────────────────────
thin = Side(style="thin")
full_border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style(cell, bold=False, red=False, header=False, num_format=None):
    cell.border = full_border
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.font = Font(bold=(bold or header), color="FF0000" if red else "000000",
                     size=12 if header else 11)
    if header:
        cell.fill = PatternFill("solid", fgColor="87CEEB")
    if num_format:
        cell.number_format = num_format

def row(ws, r, label, value, bold_label=True, num_format="0.0000"):
    c1 = ws.cell(r, 1, label)
    c2 = ws.cell(r, 2, value)
    style(c1, bold=bold_label)
    style(c2, num_format=num_format)

# ══ Sheet 1 ═══════════════════════════════════════════════════════════════════
wb = Workbook()
ws1 = wb.active
ws1.title = "Завдання 1"
ws1.column_dimensions["A"].width = 55
ws1.column_dimensions["B"].width = 20

style(ws1.cell(1, 1, "Завдання №1 — Довірчий інтервал для конверсії"), header=True)
style(ws1.cell(1, 2, ""), header=True)

style(ws1.cell(2, 1, "Параметр / Формула"), bold=True)
style(ws1.cell(2, 2, "Значення"), bold=True)

row(ws1, 3,  "Коефіцієнт конверсії (p)",               p,        num_format="0.00%")
row(ws1, 4,  "Кількість користувачів (n)",              n,        num_format="0")
row(ws1, 5,  "Рівень довіри",                           0.95,     num_format="0.00%")
row(ws1, 6,  "z-score (95%)",                           z,        num_format="0.0000")
row(ws1, 7,  "Стандартна похибка  SE = √(p·(1−p)/n)",  se_1,     num_format="0.0000")
row(ws1, 8,  "Похибка  ME = z · SE",                   me_1,     num_format="0.0000")
row(ws1, 9,  "Нижня межа ДІ  p − ME",                  ci_lower, num_format="0.0000")
row(ws1, 10, "Верхня межа ДІ  p + ME",                 ci_upper, num_format="0.0000")

style(ws1.cell(11, 1, "Висновок:"), bold=True)
style(ws1.cell(11, 2,
    f"95% ДІ: [{ci_lower*100:.2f}% ; {ci_upper*100:.2f}%]"), red=True, bold=True)

# ══ Sheet 2 ═══════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Завдання 2")
ws2.column_dimensions["A"].width = 60
ws2.column_dimensions["B"].width = 20

style(ws2.cell(1, 1, "Завдання №2 — A/B-тест: порівняння конверсій"), header=True)
style(ws2.cell(1, 2, ""), header=True)

style(ws2.cell(2, 1, "Параметр / Формула"), bold=True)
style(ws2.cell(2, 2, "Значення"), bold=True)

row(ws2, 3,  "Конверсія групи A  (p̂_A)",                                              p_a,    num_format="0.00%")
row(ws2, 4,  "Розмір групи A  (n_A)",                                                  n_a,    num_format="0")
row(ws2, 5,  "Конверсія групи B  (p̂_B)",                                              p_b,    num_format="0.00%")
row(ws2, 6,  "Розмір групи B  (n_B)",                                                  n_b,    num_format="0")
row(ws2, 7,  "Об'єднана конверсія  p_pool = (n_A·p̂_A + n_B·p̂_B) / (n_A + n_B)",     p_pool, num_format="0.0000")
row(ws2, 8,  "Стандартна похибка  SE = √(p_pool·(1−p_pool)·(1/n_A + 1/n_B))",        se_2,   num_format="0.0000")
row(ws2, 9,  "Z-статистика  z = (p̂_B − p̂_A) / SE",                                   z_stat, num_format="0.0000")
row(ws2, 10, "Критичне значення z (95%, двостороннє)",                                 z_crit, num_format="0.00")

conclusion = (
    f"Оскільки Z = {z_stat:.4f} перевищує критичне значення {z_crit}, "
    "різниця між конверсіями груп A і B є статистично значущою."
    if is_significant else
    f"Оскільки Z = {z_stat:.4f} не перевищує критичне значення {z_crit}, "
    "різниця між конверсіями груп A і B НЕ є статистично значущою."
)

style(ws2.cell(11, 1, "Висновок:"), bold=True)
style(ws2.cell(11, 2, conclusion), red=True, bold=True)

# ── Save ──────────────────────────────────────────────────────────────────────
output_path = "Moiseyenko_Alona_PA_assignment_7.xlsx"
wb.save(output_path)
print(f"Saved: {output_path}")
